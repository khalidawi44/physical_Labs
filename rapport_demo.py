# -*- coding: utf-8 -*-
"""Rapport de démonstration Alliance Groupe — le livrable, en exemple.

Ce module construit, à partir du modèle éditable `alliance_modele.py`, un
**exemple fictif** du livrable qu'Alliance Groupe remet à la fin d'un audit :
un rapport brandé (constats classés par gravité + synthèse) et un devis de
remédiation chiffré. Aucun site réel n'est audité : c'est une démonstration,
pour montrer à un prospect ce qu'il recevrait.

Deux sorties :
- des tableaux (`pandas.DataFrame`) pour l'aperçu à l'écran ;
- un document Word (`document_docx()` → octets) téléchargeable.
"""
from __future__ import annotations

from io import BytesIO
from typing import Dict, List

import pandas as pd

import alliance_modele as m

TITRE = "Rapport d'audit de sécurité — exemple de démonstration"
AVERTISSEMENT = ("Exemple fictif de livrable. Aucun site réel n'est concerné ; les constats et les "
                 "montants sont illustratifs et servent uniquement à montrer le rendu du rapport.")


def _ordre_gravite(f: Dict[str, str]) -> int:
    return m.GRAVITES.index(f["gravite"])


def findings_tries() -> List[Dict[str, str]]:
    """Les constats d'exemple, du plus grave au moins grave."""
    return sorted(m.findings_demo(), key=_ordre_gravite)


def compte_par_gravite() -> Dict[str, int]:
    """Nombre de constats par niveau de gravité (dans l'ordre officiel)."""
    compte = {g: 0 for g in m.GRAVITES}
    for f in m.findings_demo():
        compte[f["gravite"]] += 1
    return compte


def tableau_findings() -> pd.DataFrame:
    return pd.DataFrame([{"Réf.": f["id"], "Gravité": f["gravite"], "Composant": f["composant"],
                          "Outil": f["outil"], "Constat": f["constat"], "Recommandation": f["recommandation"]}
                         for f in findings_tries()])


def tableau_devis() -> pd.DataFrame:
    lignes = [{"Poste": d["poste"], "Détail": d["detail"], "Montant (€ HT)": int(d["montant"])}
              for d in m.devis_demo()]
    lignes.append({"Poste": "Total", "Détail": "", "Montant (€ HT)": m.total_devis()})
    return pd.DataFrame(lignes)


def synthese() -> str:
    c = compte_par_gravite()
    haut = c["Critique"] + c["Élevé"]
    return (f"L'audit d'exemple relève **{len(m.findings_demo())} constats**, dont **{haut} à traiter en priorité** "
            f"({c['Critique']} critique, {c['Élevé']} élevé). La remédiation chiffrée s'élève à "
            f"**{m.total_devis()} € HT** (exemple), contre-audit de validation inclus.")


def document_docx() -> bytes:
    """Construit le rapport Word brandé et le renvoie en octets (pour téléchargement)."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    titre = doc.add_heading(m.MARQUE, level=0)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st = doc.add_paragraph(TITRE)
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    st.runs[0].italic = True

    av = doc.add_paragraph()
    r = av.add_run("⚠ " + AVERTISSEMENT)
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x39, 0x2B)
    r.font.size = Pt(9)

    doc.add_paragraph(f"Client : {m.CLIENT_DEMO}")

    doc.add_heading("1. Synthèse", level=1)
    doc.add_paragraph(synthese().replace("**", ""))
    compte = compte_par_gravite()
    t = doc.add_table(rows=1, cols=2)
    t.style = "Light Grid Accent 1"
    t.rows[0].cells[0].text = "Gravité"
    t.rows[0].cells[1].text = "Nombre"
    for g in m.GRAVITES:
        c = t.add_row().cells
        c[0].text = g
        c[1].text = str(compte[g])

    doc.add_heading("2. Constats détaillés", level=1)
    for f in findings_tries():
        h = doc.add_heading(f"{f['id']} · {f['composant']} — {f['gravite']}", level=2)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)
        p = doc.add_paragraph()
        p.add_run("Outil : ").bold = True
        p.add_run(f["outil"])
        p = doc.add_paragraph()
        p.add_run("Constat : ").bold = True
        p.add_run(f["constat"])
        p = doc.add_paragraph()
        p.add_run("Recommandation : ").bold = True
        p.add_run(f["recommandation"])

    doc.add_heading("3. Devis de remédiation (exemple)", level=1)
    dt = doc.add_table(rows=1, cols=3)
    dt.style = "Light Grid Accent 1"
    for i, entete in enumerate(("Poste", "Détail", "Montant (€ HT)")):
        dt.rows[0].cells[i].text = entete
    for d in m.devis_demo():
        c = dt.add_row().cells
        c[0].text = d["poste"]
        c[1].text = d["detail"]
        c[2].text = f"{int(d['montant'])} €"
    total = dt.add_row().cells
    total[0].text = "Total"
    total[2].text = f"{m.total_devis()} € HT"
    for cell in (total[0], total[2]):
        for para in cell.paragraphs:
            for run in para.runs:
                run.bold = True

    doc.add_paragraph()
    pied = doc.add_paragraph(f"{m.MARQUE} — rapport de démonstration. Du constat à la correction chiffrée.")
    pied.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pied.runs[0].italic = True

    tampon = BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
